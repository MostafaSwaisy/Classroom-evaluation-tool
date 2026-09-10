"""P3-U5: grading_workspace (manual path) — 3 panes, autosave/restore via R11,
status chip, code tree from real extracted/, file-won't-open flag, disabled AI."""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import Workbook

from gui.screens.grading_workspace import Screen

REPO = Path(__file__).resolve().parent.parent

_HEADERS = ["الرقم الجامعي", "الاسم", "الإيميل", "الحالة", "متأخر",
            "وقت التسليم", "عدد الملفات", "الملفات", "روابط", "الدرجة الحالية"]

_STUDENTS = [(f"12021{i:04d}", f"طالب رقم {i}") for i in range(1, 15)]


def _make_roster(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "التسليمات"
    ws["A1"] = "واجب: اختبار"
    for c, h in enumerate(_HEADERS, start=1):
        ws.cell(row=5, column=c, value=h)
    for r, (sid, name) in enumerate(_STUDENTS, start=6):
        ws.cell(row=r, column=1, value=sid)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=4, value="سلّم")
        ws.cell(row=r, column=7, value=2)
    wb.save(path)


@pytest.fixture
def work_dir(tmp_path: Path) -> Path:
    wd = tmp_path / "submissions" / "PHP2026" / "HW01"
    (wd / "extracted").mkdir(parents=True)
    _make_roster(wd / "_roster.xlsx")
    # a real extracted tree for the first two students
    for sid, name in _STUDENTS[:2]:
        d = wd / "extracted" / f"{sid}_{name}"
        d.mkdir()
        (d / "Q1.php").write_text("<?php\n// حل السؤال الأول\necho 'مرحبا';\n", encoding="utf-8")
        (d / "Q2.php").write_text("<?php function f(){ return 2; }\n", encoding="utf-8")
    return wd


@pytest.fixture
def rubrics_dir(tmp_path: Path) -> Path:
    d = tmp_path / "rubrics"
    d.mkdir()
    (d / "hw01.yaml").write_text(
        "assignment: PHP2026/HW01\n"
        "max_points: 10\n"
        "criteria:\n"
        "  - key: correctness\n"
        "    label: الصحة\n"
        "    points: 6\n"
        "  - key: style\n"
        "    label: الأسلوب\n"
        "    points: 4\n",
        encoding="utf-8",
    )
    return d


class _Services:
    def __init__(self, work_dir, rubrics_dir) -> None:  # noqa: ANN001
        self.active_assignment_dir = str(work_dir)
        self.rubrics_dir = str(rubrics_dir)
        self.config_path = None
        self.backend = None


@pytest.fixture
def screen(qtbot, work_dir, rubrics_dir):
    s = Screen(_Services(work_dir, rubrics_dir))
    qtbot.addWidget(s)
    s.load()
    return s


def _select_student(screen, index: int):
    """Select the (index)-th real student item, skipping batch headers."""
    seen = 0
    for i in range(screen._list.count()):
        it = screen._list.item(i)
        if it.data(0x0100):  # Qt.UserRole -> key set only on real rows
            if seen == index:
                screen._list.setCurrentItem(it)
                return it
            seen += 1
    raise AssertionError("student not found")


# --- construction + states -----------------------------------
def test_empty_state_without_a_prepared_assignment(qtbot, rubrics_dir, tmp_path):
    s = Screen(_Services(tmp_path / "nope", rubrics_dir))
    qtbot.addWidget(s)
    s.load()
    assert s.state_view.state == "empty"


def test_loads_ok_with_three_panes_and_the_rubric(screen):
    assert screen.state_view.state == "ok"
    assert set(screen._score_inputs) == {"correctness", "style"}
    assert screen._score_inputs["correctness"].maximum() == 6


def test_error_state_on_unreadable_roster(qtbot, work_dir, rubrics_dir):
    (work_dir / "_roster.xlsx").write_bytes(b"not an xlsx")
    s = Screen(_Services(work_dir, rubrics_dir))
    qtbot.addWidget(s)
    s.load()
    assert s.state_view.state == "error"


# --- affordances -------------------------------------------
def test_student_list_is_grouped_into_batches_of_ten(screen):
    headers = [screen._list.item(i).text() for i in range(screen._list.count())
               if not screen._list.item(i).data(0x0100)]
    assert any("الدفعة 1" in h for h in headers)
    assert any("الدفعة 2" in h for h in headers)


def test_picking_a_student_loads_the_code_tree_from_extracted(screen):
    _select_student(screen, 0)
    labels = _tree_labels(screen._tree)
    assert "Q1.php" in labels and "Q2.php" in labels


def test_picking_a_file_shows_its_utf8_content_readonly(screen):
    _select_student(screen, 0)
    node = _find_node(screen._tree, "Q1.php")
    screen._tree.setCurrentItem(node)
    assert "مرحبا" in screen._code.toPlainText()
    assert screen._code.isReadOnly()


def test_scores_update_the_running_total(screen):
    _select_student(screen, 0)
    screen._score_inputs["correctness"].setValue(5)
    screen._score_inputs["style"].setValue(3)
    assert "8" in screen._total_label.text()


def test_status_flips_to_draft_on_first_edit(screen):
    it = _select_student(screen, 0)
    assert "لم يبدأ" in it.text()
    screen._score_inputs["style"].setValue(2)
    assert "مسودة" in it.text()


def test_file_wont_open_nulls_scores_and_adds_the_flag(screen):
    _select_student(screen, 0)
    screen._on_file_wont_open()
    key = screen._current_key
    entry = screen._state.get_entry(key)
    assert entry["scores"] is None
    assert "الملف ما بينفتح" in entry["flags"]


def test_similarity_flag_is_recorded_with_the_id(screen):
    _select_student(screen, 0)
    screen._similar_id.setText("120210999")
    entry = screen._state.get_entry(screen._current_key)
    assert "تشابه مع 120210999" in entry["flags"]


def test_ai_panel_is_present_but_disabled(screen):
    from PySide6.QtWidgets import QGroupBox
    ai = next(g for g in screen.findChildren(QGroupBox) if g.title().startswith("مساعدة AI"))
    assert ai.isEnabled() is False
    assert "اربط Claude" in screen._ai_link.text()


def test_close_and_reopen_restores_every_entry(qtbot, work_dir, rubrics_dir):
    s1 = Screen(_Services(work_dir, rubrics_dir))
    qtbot.addWidget(s1)
    s1.load()
    _select_student(s1, 0)
    s1._score_inputs["correctness"].setValue(4)
    s1._feedback.setPlainText("ملاحظة تجريبية")
    s1._state.flush()

    s2 = Screen(_Services(work_dir, rubrics_dir))
    qtbot.addWidget(s2)
    s2.load()
    it = _select_student(s2, 0)
    assert "مسودة" in it.text()
    assert s2._score_inputs["correctness"].value() == 4
    assert s2._feedback.toPlainText() == "ملاحظة تجريبية"


def test_state_sidecar_is_the_json_only(screen, work_dir):
    _select_student(screen, 0)
    screen._feedback.setPlainText("x")
    screen._state.flush()
    assert (work_dir / "_grading_state.json").exists()
    assert not (work_dir / "_grading_state.md").exists()


# --- guardrail --------------------------------------------
def test_no_upload_tokens_in_source():
    src = (REPO / "gui" / "screens" / "grading_workspace.py").read_text(encoding="utf-8")
    assert re.search(r"\b(push|upload|confirm|sync)\b|--confirm", src, re.IGNORECASE) is None


# --- helpers ----------------------------------------------
def _tree_labels(tree) -> list[str]:
    out = []

    def walk(item):
        for i in range(item.childCount()):
            ch = item.child(i)
            out.append(ch.text(0))
            walk(ch)

    walk(tree.invisibleRootItem())
    return out


def _find_node(tree, name: str):
    def walk(item):
        for i in range(item.childCount()):
            ch = item.child(i)
            if ch.text(0) == name:
                return ch
            found = walk(ch)
            if found:
                return found
        return None

    return walk(tree.invisibleRootItem())
