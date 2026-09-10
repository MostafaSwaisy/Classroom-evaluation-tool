"""P3-U6: grades_draft — editable table, in-app stats, persistent draft banner,
Export via worker → write_grades."""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QObject, Signal

from gui.screens.grades_draft import _JOB_EXPORT, Screen

REPO = Path(__file__).resolve().parent.parent

_HEADERS = ["الرقم الجامعي", "الاسم", "الإيميل", "الحالة", "متأخر",
            "وقت التسليم", "عدد الملفات", "الملفات", "روابط", "الدرجة الحالية"]


def _make_roster(path: Path, rows: list[tuple[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "التسليمات"
    for c, h in enumerate(_HEADERS, start=1):
        ws.cell(row=5, column=c, value=h)
    for r, (sid, name) in enumerate(rows, start=6):
        ws.cell(row=r, column=1, value=sid)
        ws.cell(row=r, column=2, value=name)
    wb.save(path)


_STATE = {
    "12021001": {"scores": {"correctness": 5, "style": 3}, "feedback": "جيد جداً",
                 "flags": [], "status": "مسودة"},
    "12021002": {"scores": {"correctness": 6, "style": 4}, "feedback": "",
                 "flags": ["تشابه مع 12021001"], "status": "مسودة"},
    "12021003": {"scores": None, "feedback": "ملف تالف",
                 "flags": ["الملف ما بينفتح"], "status": "مسودة"},
}


class _FakeWorker(QObject):
    finished = Signal(str, object)
    failed = Signal(str, str, str, str)


class _FakeBackend:
    def __init__(self) -> None:
        self.worker = _FakeWorker()
        self.jobs: list[tuple] = []

    def submit(self, job_id: str, fn) -> None:  # noqa: ANN001
        self.jobs.append((job_id, fn))


class _Services:
    def __init__(self, work_dir, rubrics_dir) -> None:  # noqa: ANN001
        self.active_assignment_dir = str(work_dir)
        self.rubrics_dir = str(rubrics_dir)
        self.config_path = None
        self.backend = _FakeBackend()


@pytest.fixture
def rubrics_dir(tmp_path: Path) -> Path:
    d = tmp_path / "rubrics"
    d.mkdir()
    (d / "hw.yaml").write_text(
        "assignment: PHP2026/HW01\nmax_points: 10\ncriteria:\n"
        "  - key: correctness\n    label: الصحة\n    points: 6\n"
        "  - key: style\n    label: الأسلوب\n    points: 4\n",
        encoding="utf-8")
    return d


@pytest.fixture
def work_dir(tmp_path: Path) -> Path:
    wd = tmp_path / "submissions" / "PHP2026" / "HW01"
    wd.mkdir(parents=True)
    _make_roster(wd / "_roster.xlsx",
                 [("12021001", "طالب أول"), ("12021002", "طالب ثاني"),
                  ("12021003", "طالب ثالث")])
    (wd / "_grading_state.json").write_text(
        json.dumps(_STATE, ensure_ascii=False), encoding="utf-8")
    return wd


@pytest.fixture
def screen(qtbot, work_dir, rubrics_dir):
    s = Screen(_Services(work_dir, rubrics_dir))
    qtbot.addWidget(s)
    s.load()
    return s


# --- states --------------------------------------------------
def test_empty_without_a_grading_state_file(qtbot, tmp_path, rubrics_dir):
    wd = tmp_path / "empty"
    wd.mkdir()
    s = Screen(_Services(wd, rubrics_dir))
    qtbot.addWidget(s)
    s.load()
    assert s.state_view.state == "empty"


def test_loads_ok_with_a_row_per_graded_student(screen):
    assert screen.state_view.state == "ok"
    assert screen._table.rowCount() == 3
    headers = [screen._table.horizontalHeaderItem(c).text()
               for c in range(screen._table.columnCount())]
    assert headers[:2] == ["الرقم الجامعي", "الاسم"]
    assert "الصحة" in headers and "الأسلوب" in headers
    assert "المجموع" in headers


def test_error_state_on_unreadable_roster(qtbot, work_dir, rubrics_dir):
    (work_dir / "_roster.xlsx").write_bytes(b"nope")
    s = Screen(_Services(work_dir, rubrics_dir))
    qtbot.addWidget(s)
    s.load()
    assert s.state_view.state == "error"


# --- table --------------------------------------------------
def test_total_column_is_the_sum_of_criteria(screen):
    total_col = 2 + 2  # id, name, 2 criteria
    assert screen._table.item(0, total_col).text() == "8"
    assert screen._table.item(1, total_col).text() == "10"
    assert screen._table.item(2, total_col).text() == "—"   # null scores


def test_flagged_and_null_rows_are_painted_differently(screen):
    normal = screen._table.item(0, 0).background().color().name()
    flagged = screen._table.item(1, 0).background().color().name()
    nulled = screen._table.item(2, 0).background().color().name()
    assert flagged != normal
    assert nulled != normal
    assert flagged != nulled


def test_editing_a_score_updates_total_and_persists(screen, work_dir):
    screen._table.item(0, 2).setText("2")      # correctness 5 -> 2
    assert screen._table.item(0, 4).text() == "5"   # 2 + 3 (total col = 2 + nc)
    screen._state.flush()
    reloaded = json.loads((work_dir / "_grading_state.json").read_text(encoding="utf-8"))
    assert reloaded["12021001"]["scores"]["correctness"] == 2


def test_editing_flags_repaints_the_row(screen):
    normal = screen._table.item(0, 0).background().color().name()
    screen._table.item(0, 6).setText("يحتاج مراجعة شفوية")   # flags col = 4 + nc
    assert screen._table.item(0, 0).background().color().name() != normal


# --- stats --------------------------------------------------
def test_stats_panel_is_computed_in_app(screen):
    assert screen._stat_labels["count"].text() == "2"        # 2 non-null totals
    assert screen._stat_labels["avg"].text() == "9"          # (8 + 10) / 2
    assert screen._stat_labels["max"].text() == "10"
    assert screen._stat_labels["min"].text() == "8"
    assert screen._stat_labels["below"].text() == "0"        # none < 5


# --- banner -----------------------------------------------
def test_draft_banner_is_present_and_has_no_close_handler():
    src = (REPO / "gui" / "screens" / "grades_draft.py").read_text(encoding="utf-8")
    assert "مسودة — لم تُرفع" in src
    assert not re.search(r"dismiss|close_banner|_banner.*hide|removeWidget.*_banner", src)


# --- export ----------------------------------------------
def test_export_submits_the_worker_job(screen):
    screen._on_export()
    assert screen.services.backend.jobs
    assert screen.services.backend.jobs[-1][0] == _JOB_EXPORT


def test_export_job_produces_a_valid_xlsx(screen, work_dir):
    screen._on_export()
    _job_id, fn = screen.services.backend.jobs[-1]
    result = fn(None)
    out = Path(result["path"])
    assert out.exists() and out.name == "grades_draft.xlsx"
    wb = load_workbook(out)
    assert wb.sheetnames == ["الدرجات", "إحصائيات"]


def test_finished_signal_shows_the_path_and_open_folder(screen, work_dir):
    screen._on_export()
    screen.services.backend.worker.finished.emit(
        _JOB_EXPORT, {"path": str(work_dir / "grades_draft.xlsx")})
    assert "grades_draft.xlsx" in screen._result_label.text()
    assert not screen._open_folder_btn.isHidden()


# --- guardrail ------------------------------------------
def test_no_upload_or_classroom_tokens_in_source():
    src = (REPO / "gui" / "screens" / "grades_draft.py").read_text(encoding="utf-8")
    assert re.search(r"\b(push|upload|confirm|sync)\b|--confirm", src, re.IGNORECASE) is None
    assert "classroom" not in src.lower()
