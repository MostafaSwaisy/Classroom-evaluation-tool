"""P1-U4 (R5): compute_status() returns structured data; the xlsx writer consumes it."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from classroom_tool import status as st

REPO = Path(__file__).resolve().parent.parent
GOLDEN_XLSX = REPO / "tests" / "golden" / "status_stubbed.xlsx"

_WORKS = [
    {"id": "w1", "title": "HW01 Routing", "workType": "ASSIGNMENT"},
    {"id": "w2", "title": "HW02 Controllers", "workType": "ASSIGNMENT"},
]
_SUBS = {
    "w1": [
        {"userId": "u_A", "state": "TURNED_IN", "late": False, "assignedGrade": 8},
        {"userId": "u_B", "state": "CREATED", "late": False, "assignedGrade": None},
    ],
    "w2": [
        {"userId": "u_A", "state": "TURNED_IN", "late": True, "assignedGrade": None},
    ],
}
_STUDENTS = {
    "u_A": {"student_id": "120210001", "full_name": "أحمد علي"},
    "u_B": {"student_id": None, "full_name": "طالب بلا رقم"},
}


@pytest.fixture(autouse=True)
def _stub_api(monkeypatch):
    monkeypatch.setattr(st.api, "list_coursework", lambda c, cid: list(_WORKS))
    monkeypatch.setattr(st.api, "list_submissions", lambda c, cid, wid: list(_SUBS.get(wid, [])))
    monkeypatch.setattr(st, "build_student_index", lambda c, cid, pat: dict(_STUDENTS))


def _compute(threshold=0.6):
    return st.compute_status(object(), {"student_id_pattern": r"^(\d+)@"}, "999",
                             risk_threshold=threshold)


def test_matrix_has_a_cell_per_submitted_pair():
    m = _compute()["matrix"]
    assert m["u_A"]["w1"] == {"submitted": True, "late": False, "grade": 8}
    assert m["u_A"]["w2"] == {"submitted": True, "late": True, "grade": None}
    assert m["u_B"]["w1"]["submitted"] is False


def test_rows_are_sorted_and_carry_the_trailing_metrics():
    rows = _compute()["rows"]
    assert [r["student_id"] for r in rows] == ["120210001", ""]  # noid sorts last, blanked

    a, b = rows
    assert a["name"] == "أحمد علي"
    assert a["ratio"] == 1.0
    assert a["late"] == 1
    assert a["avg"] == 8.0
    assert a["status"] == "جيد"

    assert b["ratio"] == 0.0
    assert b["avg"] is None
    assert b["status"] == "⚠️ متابعة"


def test_summary_and_threshold_flow_into_at_risk():
    at_06 = _compute(0.6)["summary"]
    assert at_06["total_students"] == 2
    assert at_06["total_works"] == 2
    assert at_06["threshold"] == 0.6
    assert [r["student_id"] for r in at_06["at_risk"]] == [""]      # u_B, ratio 0.0 < 0.6

    at_00 = _compute(0.0)["summary"]
    assert at_00["at_risk"] == []                                   # 0.0 < 0.0 is False


def test_status_xlsx_is_cell_for_cell_unchanged(tmp_path):
    cfg = {"output_dir": str(tmp_path), "student_id_pattern": r"^(\d+)@"}
    out = st.status(object(), cfg, "SE2026", "999", risk_threshold=0.6)

    got = load_workbook(out)["المتابعة"]
    want = load_workbook(GOLDEN_XLSX)["المتابعة"]
    assert got.max_row == want.max_row and got.max_column == want.max_column
    for row in range(1, want.max_row + 1):
        for col in range(1, want.max_column + 1):
            if (row, col) == (2, 1):
                continue  # A2 carries a datetime.now() stamp
            assert got.cell(row, col).value == want.cell(row, col).value, f"cell {row},{col}"
