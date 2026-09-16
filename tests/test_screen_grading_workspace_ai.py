"""P4-U6: grading_workspace AI panel live — 5 §5.11 states, explicit accept,
review-gated «مكتمل», manual editing unaffected."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import Workbook
from PySide6.QtCore import QObject, Signal

import gui.screens.grading_workspace as gw
from gui.screens.grading_workspace import _JOB_AI, Screen

_HEADERS = ["الرقم الجامعي", "الاسم", "الإيميل", "الحالة", "متأخر",
            "وقت التسليم", "عدد الملفات", "الملفات", "روابط", "الدرجة الحالية"]


def _roster(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "التسليمات"
    for c, h in enumerate(_HEADERS, start=1):
        ws.cell(row=5, column=c, value=h)
    for r, (sid, name) in enumerate([("12021001", "طالب أول"),
                                     ("12021002", "طالب ثاني")], start=6):
        ws.cell(row=r, column=1, value=sid)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=4, value="سلّم")
    wb.save(path)


class _FakeWorker(QObject):
    finished = Signal(str, object)
    failed = Signal(str, str, str, str)
    progress = Signal(str, int, int)
    cancelled = Signal(str)


class _FakeBackend:
    def __init__(self):
        self.worker = _FakeWorker()
        self.jobs = []
        self.cancelled = False

    def submit(self, job_id, fn):  # noqa: ANN001
        self.jobs.append((job_id, fn))

    def cancel(self):
        self.cancelled = True


@dataclass
class _St:
    provider: str = "claude_cli"
    state: str = "ready"
    detail: str = ""


class _Services:
    def __init__(self, wd, rub):
        self.active_assignment_dir = str(wd)
        self.rubrics_dir = str(rub)
        self.config_path = None
        self.backend = _FakeBackend()


@pytest.fixture
def env(tmp_path):
    wd = tmp_path / "submissions" / "PHP2026" / "HW01"
    (wd / "extracted" / "12021001_طالب أول").mkdir(parents=True)
    (wd / "extracted" / "12021001_طالب أول" / "Q1.php").write_text(
        "<?php echo 1;", encoding="utf-8")
    _roster(wd / "_roster.xlsx")
    rub = tmp_path / "rubrics"
    rub.mkdir()
    (rub / "hw.yaml").write_text(
        "assignment: PHP2026/HW01\nmax_points: 10\ncriteria:\n"
        "  - key: correctness\n    label: الصحة\n    points: 6\n"
        "  - key: style\n    label: الأسلوب\n    points: 4\n", encoding="utf-8")
    return wd, rub


@pytest.fixture
def screen(qtbot, env, monkeypatch):
    monkeypatch.setattr(gw, "_safe_status", lambda _cfg: "ready")
    wd, rub = env
    s = Screen(_Services(wd, rub))
    qtbot.addWidget(s)
    s.load()
    _pick_first(s)
    return s


def _pick_first(s):
    for i in range(s._list.count()):
        it = s._list.item(i)
        if it.data(0x0100):
            s._list.setCurrentItem(it)
            return


def _ai_state(s) -> str:
    return gw._AI_STATES[s._ai_stack.currentIndex()]


def test_clicking_suggest_with_no_student_selected_shows_an_error_not_silence(env, qtbot):
    """was a silent no-op before -- clicking اقترح with nothing selected in
    the student list did nothing at all: no error, no state change, no clue
    why it "didn't work". Confirmed against mostafa's screenshot: the panel
    was stuck on its idle page with no feedback."""
    wd, rub = env
    s = Screen(_Services(wd, rub))
    qtbot.addWidget(s)
    assert s._current_key is None
    s._run_ai(whole_batch=False)
    assert _ai_state(s) == "error"
    assert s.services.backend.jobs == []


# --- the 5 §5.11 states -------------------------------------
def test_state_not_connected(screen, monkeypatch):
    monkeypatch.setattr(gw, "_safe_status", lambda _c: "not_installed")
    screen._run_ai(whole_batch=False)
    assert _ai_state(screen) == "not_connected"
    assert screen.services.backend.jobs == []


def test_state_not_logged_in(screen, monkeypatch):
    monkeypatch.setattr(gw, "_safe_status", lambda _c: "not_logged_in")
    screen._run_ai(whole_batch=False)
    assert _ai_state(screen) == "not_logged_in"


def test_state_running_then_suggestions_shown(screen):
    screen._run_ai(whole_batch=False)
    assert _ai_state(screen) == "running"
    job_id, _fn = screen.services.backend.jobs[-1]
    assert job_id == _JOB_AI
    screen.services.backend.worker.finished.emit(_JOB_AI, {
        screen._current_key: {"scores": {"correctness": 5, "style": 3},
                              "feedback": "ملاحظة مقترحة", "flags": [], "error": None},
    })
    assert _ai_state(screen) == "shown"


def test_batch_finish_shows_a_pass_fail_summary_without_clicking_through_students(screen):
    """mostafa's report: after a whole-batch run, the only way to tell which
    of the 10 students actually got a usable suggestion was to click through
    each one individually -- there was no aggregate result anywhere."""
    screen._run_ai(whole_batch=True)
    screen.services.backend.worker.finished.emit(_JOB_AI, {
        screen._current_key: {"scores": {"correctness": 5, "style": 3},
                              "feedback": "ok", "flags": [], "error": None},
        "12021002": {"scores": {}, "feedback": "", "flags": [],
                    "error": "ردّ غير صالح"},
    })
    assert not screen._ai_batch_summary.isHidden()
    text = screen._ai_batch_summary.text()
    assert "1 نجح" in text
    assert "1 فشل" in text


def test_single_student_suggest_never_shows_the_batch_summary(screen):
    screen._run_ai(whole_batch=False)
    screen.services.backend.worker.finished.emit(_JOB_AI, {
        screen._current_key: {"scores": {"correctness": 5, "style": 3},
                              "feedback": "ok", "flags": [], "error": None},
    })
    assert screen._ai_batch_summary.isHidden()


def test_batch_progress_updates_the_running_bar(screen):
    """mostafa's report: no visible sign a batch suggest run is actually
    progressing. ai_suggest_job already emits ctx.progress(key, done, total)
    per student (gui/grading_ai.py) -- the panel just never rendered it."""
    screen._run_ai(whole_batch=True)
    assert _ai_state(screen) == "running"
    screen.services.backend.worker.progress.emit("مقترح: 12021001", 1, 2)
    assert screen._ai_progress._bar.maximum() == 2
    assert screen._ai_progress._bar.value() == 1
    assert screen._ai_progress._counter.text() == "1 من 2"


def test_state_error_on_worker_failure(screen):
    screen._run_ai(whole_batch=False)
    screen.services.backend.worker.failed.emit(_JOB_AI, "RuntimeError", "boom", "")
    assert _ai_state(screen) == "error"


def test_state_error_when_suggestion_carries_an_error(screen):
    screen._run_ai(whole_batch=False)
    screen.services.backend.worker.finished.emit(_JOB_AI, {
        screen._current_key: {"scores": {}, "feedback": "", "flags": [],
                              "error": "ردّ غير صالح"}})
    assert _ai_state(screen) == "error"


# --- nothing applied without an explicit accept -----------
def test_suggestion_does_not_touch_scores_until_accepted(screen):
    screen._run_ai(whole_batch=False)
    screen.services.backend.worker.finished.emit(_JOB_AI, {
        screen._current_key: {"scores": {"correctness": 5, "style": 3},
                              "feedback": "س", "flags": [], "error": None}})
    assert screen._score_inputs["correctness"].value() == 0  # unchanged

    screen._accept_all()
    assert screen._score_inputs["correctness"].value() == 5
    assert screen._score_inputs["style"].value() == 3
    assert screen._feedback.toPlainText() == "س"


# --- «مكتمل» is review-gated -----------------------------
def test_status_only_reaches_complete_after_the_review_checkbox(screen):
    screen._score_inputs["correctness"].setValue(4)
    entry = screen._state.get_entry(screen._current_key)
    assert entry["status"] == "مسودة"

    screen._reviewed_cb.setChecked(True)
    entry = screen._state.get_entry(screen._current_key)
    assert entry["status"] == "مكتمل"


# --- manual editing unaffected with no `claude` ----------
def test_manual_editing_still_works_when_ai_not_connected(env, qtbot, monkeypatch):
    monkeypatch.setattr(gw, "_safe_status", lambda _c: "not_installed")
    wd, rub = env
    s = Screen(_Services(wd, rub))
    qtbot.addWidget(s)
    s.load()
    _pick_first(s)
    s._score_inputs["correctness"].setValue(6)
    s._feedback.setPlainText("تصحيح يدوي")
    entry = s._state.get_entry(s._current_key)
    assert entry["scores"]["correctness"] == 6
    assert entry["feedback"] == "تصحيح يدوي"


# --- guardrail --------------------------------------------
def test_no_upload_tokens_in_source():
    import re
    src = (Path(__file__).resolve().parent.parent
           / "gui" / "screens" / "grading_workspace.py").read_text(encoding="utf-8")
    assert re.search(r"\b(push|upload|confirm|sync)\b|--confirm", src, re.IGNORECASE) is None


# --- the AI batch is the longest job in the app; it needs a cancel ------
def test_running_panel_is_the_shared_progress_panel_and_is_cancellable(screen):
    screen._run_ai(whole_batch=True)
    assert _ai_state(screen) == "running"
    assert screen._ai_progress.is_running
    assert not screen._ai_progress._cancel_btn.isHidden()


def test_ticks_drive_a_percentage_and_a_count(screen):
    screen._run_ai(whole_batch=True)
    screen.services.backend.worker.progress.emit("مقترح: 120210123", 1, 4)
    assert screen._ai_progress._percent.text() == "25%"
    assert screen._ai_progress._counter.text() == "1 من 4"
    assert "120210123" in screen._ai_progress._detail.text()


def test_ticks_ignored_when_the_ai_panel_is_not_running(screen):
    screen.services.backend.worker.progress.emit("شيء تاني", 2, 8)
    assert screen._ai_progress._percent.text() == ""


def test_cancelling_the_batch_asks_the_backend_to_stop(screen):
    screen._run_ai(whole_batch=True)
    screen._ai_progress._cancel_btn.click()
    assert screen.services.backend.cancelled


def test_cancelled_batch_stops_the_panel_and_says_what_was_kept(screen):
    screen._run_ai(whole_batch=True)
    screen.services.backend.worker.cancelled.emit(gw._JOB_AI)
    assert not screen._ai_progress.is_running
    assert _ai_state(screen) == "error"
    assert "أُلغي" in screen._ai_msgs["error"].text()


def test_finishing_stops_the_panel(screen):
    screen._run_ai(whole_batch=False)
    screen.services.backend.worker.finished.emit(gw._JOB_AI, {})
    assert not screen._ai_progress.is_running


def test_a_restart_resets_a_stale_percentage(screen):
    screen._run_ai(whole_batch=True)
    screen.services.backend.worker.progress.emit("مقترح: x", 4, 4)
    screen.services.backend.worker.finished.emit(gw._JOB_AI, {})
    screen._run_ai(whole_batch=True)
    assert screen._ai_progress._percent.text() == ""
