"""P3-U1 (R1): tools/write_grades exposes write_grades(work_dir, data) -> Path;
main() is a thin sys.argv wrapper and its stdout is unchanged."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.write_grades import write_grades

REPO = Path(__file__).resolve().parent.parent
FIXTURE_WORKDIR = REPO / "submissions" / "860473355891" / "واجب_1"
FIXTURE_JSON = FIXTURE_WORKDIR / "grades.json"


@pytest.fixture
def work_dir(tmp_path: Path) -> Path:
    d = tmp_path / "w"
    d.mkdir()
    shutil.copy2(FIXTURE_WORKDIR / "_roster.xlsx", d / "_roster.xlsx")
    return d


@pytest.fixture
def data() -> dict:
    return json.loads(FIXTURE_JSON.read_text(encoding="utf-8"))


def test_returns_a_path_that_exists(work_dir, data):
    out = write_grades(work_dir, data)
    assert isinstance(out, Path)
    assert out.exists()
    assert out == work_dir / "grades_draft.xlsx"


def test_workbook_has_both_rtl_sheets(work_dir, data):
    out = write_grades(work_dir, data)
    wb = load_workbook(out)
    assert wb.sheetnames == ["الدرجات", "إحصائيات"]
    for name in wb.sheetnames:
        assert wb[name].sheet_view.rightToLeft is True


def test_total_column_holds_a_sum_formula(work_dir, data):
    out = write_grades(work_dir, data)
    ws = load_workbook(out)["الدرجات"]
    n_crit = len(data["criteria"])
    total_col = 3 + n_crit  # id, name, then one col per criterion
    first_student_row = 6   # header at row 5
    cell = ws.cell(row=first_student_row, column=total_col)
    assert isinstance(cell.value, str) and cell.value.startswith("=SUM(")


def test_pure_function_does_not_print(work_dir, data, capsys):
    write_grades(work_dir, data)
    assert capsys.readouterr().out == ""


def test_cli_stdout_is_unchanged(work_dir):
    """main() still prints the ✓ path / summary / review lines to stdout."""
    root = Path(tempfile.mkdtemp())
    try:
        wd = root / "w"
        wd.mkdir(parents=True)
        shutil.copy2(FIXTURE_WORKDIR / "_roster.xlsx", wd / "_roster.xlsx")
        proc = subprocess.run(
            [sys.executable, str(REPO / "tools" / "write_grades.py"),
             str(wd), str(FIXTURE_JSON)],
            capture_output=True, text=True, encoding="utf-8", errors="replace",
            timeout=120, cwd=REPO, env=dict(os.environ),
        )
        assert proc.returncode == 0, proc.stderr
        out = proc.stdout.replace("\r\n", "\n")
        data = json.loads(FIXTURE_JSON.read_text(encoding="utf-8"))
        expected = (
            f"✓ {wd / 'grades_draft.xlsx'}\n"
            f"  {len(data['grades'])} طالب  |  {len(data['criteria'])} معيار  |  "
            f"العلامة الكاملة {data['max_points']}\n"
            f"  راجع الملف قبل الرفع.\n"
        )
        assert out == expected
    finally:
        shutil.rmtree(root, ignore_errors=True)
