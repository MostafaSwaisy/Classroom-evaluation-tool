"""P3-U7: end-to-end gate (spec §8 criterion 3, minus AI).

Drives the manual grading path to Export and proves the GUI's
`grades_draft.xlsx` is cell-for-cell identical (values *and* formulas, both
sheets) to a real CLI run of `tools/write_grades.py` from the same grades.json.

The three metadata rows (title / timestamp / note) are intentionally excluded
— row 2 carries `datetime.now()` and can never be byte-stable between two runs.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from gui.screens.grades_draft import _JOB_EXPORT, Screen

REPO = Path(__file__).resolve().parent.parent
FIXTURE = REPO / "submissions" / "860473355891" / "واجب_1"
GRADES_JSON = FIXTURE / "grades.json"

_ROSTER_HEADERS = ["الرقم الجامعي", "الاسم", "الإيميل", "الحالة", "متأخر",
                   "وقت التسليم", "عدد الملفات", "الملفات", "روابط", "الدرجة الحالية"]


class _FakeWorker:
    class _Sig:
        def connect(self, *_a):  # noqa: ANN001
            pass
    finished = _Sig()
    failed = _Sig()


class _FakeBackend:
    def __init__(self) -> None:
        self.worker = _FakeWorker()
        self.jobs: list = []

    def submit(self, job_id: str, fn) -> None:  # noqa: ANN001
        self.jobs.append((job_id, fn))


class _Services:
    def __init__(self, work_dir, rubrics_dir) -> None:  # noqa: ANN001
        self.active_assignment_dir = str(work_dir)
        self.rubrics_dir = str(rubrics_dir)
        self.config_path = None
        self.backend = _FakeBackend()


def _roster(path: Path, rows: list[tuple[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "التسليمات"
    for c, h in enumerate(_ROSTER_HEADERS, start=1):
        ws.cell(row=5, column=c, value=h)
    for r, (sid, name) in enumerate(rows, start=6):
        ws.cell(row=r, column=1, value=sid)
        ws.cell(row=r, column=2, value=name)
    wb.save(path)


def _cells(path: Path, sheet: str, min_row: int) -> list:
    ws = load_workbook(path)[sheet]
    return [[c.value for c in row]
            for row in ws.iter_rows(min_row=min_row)]


@pytest.fixture
def data() -> dict:
    return json.loads(GRADES_JSON.read_text(encoding="utf-8"))


def test_gui_export_matches_a_cli_run_cell_for_cell(qtbot, tmp_path, data):
    # --- CLI side: the genuine tools/write_grades.py on the fixture json ------
    dir_a = tmp_path / "cli" / "PHP2026" / "HW01"
    dir_a.mkdir(parents=True)
    _roster(dir_a / "_roster.xlsx", [("x", "y")])   # names come from the json itself
    proc = subprocess.run(
        [sys.executable, str(REPO / "tools" / "write_grades.py"),
         str(dir_a), str(GRADES_JSON)],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
        timeout=120, cwd=REPO, env=dict(os.environ),
    )
    assert proc.returncode == 0, proc.stderr
    xlsx_a = dir_a / "grades_draft.xlsx"

    # --- GUI side: seed state + rubric so the screen reproduces the same data -
    dir_b = tmp_path / "gui" / "PHP2026" / "HW01"
    dir_b.mkdir(parents=True)
    _roster(dir_b / "_roster.xlsx",
            [(g["student_id"], g["name"]) for g in data["grades"]])
    state = {
        g["student_id"]: {"scores": g["scores"], "feedback": g["feedback"],
                          "flags": g["flags"], "status": "مسودة"}
        for g in data["grades"]
    }
    (dir_b / "_grading_state.json").write_text(
        json.dumps(state, ensure_ascii=False), encoding="utf-8")

    rub = tmp_path / "rubrics"
    rub.mkdir()
    lines = ["assignment: PHP2026/HW01", f"max_points: {data['max_points']}",
             "criteria:"]
    for c in data["criteria"]:
        lines += [f"  - key: {c['key']}",
                  f"    label: {c['label']}",
                  f"    points: {c['points']}"]
    (rub / "hw.yaml").write_text("\n".join(lines) + "\n", encoding="utf-8")

    app_screen = Screen(_Services(dir_b, rub))
    qtbot.addWidget(app_screen)
    app_screen.load()
    assert app_screen.state_view.state == "ok"
    app_screen._on_export()
    _job_id, fn = app_screen.services.backend.jobs[-1]
    assert _job_id == _JOB_EXPORT
    result = fn(None)
    xlsx_b = Path(result["path"])

    # --- compare: grades grid from the header row down, full stats sheet -----
    assert _cells(xlsx_a, "الدرجات", 5) == _cells(xlsx_b, "الدرجات", 5)
    assert _cells(xlsx_a, "إحصائيات", 1) == _cells(xlsx_b, "إحصائيات", 1)
