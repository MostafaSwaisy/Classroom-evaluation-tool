#!/usr/bin/env python3
"""
يحوّل ملف درجات JSON إلى grades_draft.xlsx جاهز للمراجعة.

الاستخدام:
    python tools/write_grades.py <مجلد_الواجب> <ملف_json>

صيغة الـ JSON المتوقعة:
{
  "assignment": "HW03 - Laravel Routing",
  "max_points": 10,
  "criteria": [
    {"key": "routes",  "label": "تعريف المسارات",  "points": 3},
    {"key": "control", "label": "الـ Controllers",  "points": 4},
    {"key": "style",   "label": "جودة الكود",       "points": 3}
  ],
  "grades": [
    {
      "student_id": "120210123",
      "scores": {"routes": 3, "control": 3.5, "style": 2},
      "feedback": "المسارات ممتازة. الـ controller فيه منطق أعمال لازم ينتقل لـ service.",
      "flags": ["تشابه مع 120210456"]
    }
  ]
}
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def write_grades(work_dir: Path, data: dict) -> Path:
    """يبني grades_draft.xlsx من قاموس درجات ويعيد مسار الملف المكتوب.

    دالة نقيّة: لا تطبع شيئاً (الـ CLI يتكفّل بالإخراج، والـ GUI يشغّلها في worker).
    """
    work_dir = Path(work_dir)
    criteria = data["criteria"]
    total_possible = sum(c["points"] for c in criteria)
    max_points = data.get("max_points", total_possible)

    # اسحب أسماء الطلاب من _roster.xlsx إن وُجد
    names = _load_names(work_dir / "_roster.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "الدرجات"
    ws.sheet_view.rightToLeft = True

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="2F5597")
    total_fill = PatternFill("solid", fgColor="D9E1F2")
    flag_fill = PatternFill("solid", fgColor="FFC7CE")

    ws["A1"] = f"مسودة درجات — {data.get('assignment', '')}"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws["A2"] = (f"العلامة الكاملة: {max_points}    |    "
                f"عدد الطلاب: {len(data['grades'])}    |    "
                f"{datetime.now():%Y-%m-%d %H:%M}")
    ws["A2"].font = Font(name="Arial", italic=True, size=9)
    ws["A3"] = ("⚠️ مسودة — راجعها وعدّل الأعمدة قبل الرفع. "
                "عمود المجموع محسوب بمعادلة.")
    ws["A3"].font = Font(name="Arial", bold=True, size=9, color="C00000")

    start = 5
    headers = (["الرقم الجامعي", "الاسم"]
               + [f"{c['label']} ({c['points']})" for c in criteria]
               + ["المجموع", "ملاحظات", "تنبيهات"])
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=start, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[start].height = 32

    first_crit = 3
    last_crit = first_crit + len(criteria) - 1
    total_col = last_crit + 1

    for offset, row in enumerate(data["grades"], start=1):
        r = start + offset
        sid = str(row.get("student_id", ""))
        ws.cell(row=r, column=1, value=sid)
        ws.cell(row=r, column=2, value=row.get("name") or names.get(sid, ""))

        for i, crit in enumerate(criteria):
            value = row.get("scores", {}).get(crit["key"])
            ws.cell(row=r, column=first_crit + i, value=value)

        rng = (f"{get_column_letter(first_crit)}{r}:"
               f"{get_column_letter(last_crit)}{r}")
        total = ws.cell(row=r, column=total_col, value=f"=SUM({rng})")
        total.fill = total_fill
        total.font = Font(name="Arial", bold=True, size=10)

        ws.cell(row=r, column=total_col + 1, value=row.get("feedback", ""))
        flags = row.get("flags") or []
        flag_cell = ws.cell(row=r, column=total_col + 2, value=" | ".join(flags))
        if flags:
            flag_cell.fill = flag_fill

        for col in range(1, total_col + 3):
            cell = ws.cell(row=r, column=col)
            if not cell.font.bold:
                cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(col > total_col))

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 26
    for i in range(len(criteria)):
        ws.column_dimensions[get_column_letter(first_crit + i)].width = 14
    ws.column_dimensions[get_column_letter(total_col)].width = 10
    ws.column_dimensions[get_column_letter(total_col + 1)].width = 60
    ws.column_dimensions[get_column_letter(total_col + 2)].width = 24

    last_row = start + len(data["grades"])
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(total_col + 2)}{last_row}"

    stats = wb.create_sheet("إحصائيات")
    stats.sheet_view.rightToLeft = True
    tc = get_column_letter(total_col)
    rows = [
        ("عدد المصححين", f"=COUNT(الدرجات!{tc}{start + 1}:{tc}{last_row})"),
        ("المتوسط", f"=IFERROR(AVERAGE(الدرجات!{tc}{start + 1}:{tc}{last_row}),0)"),
        ("الأعلى", f"=IFERROR(MAX(الدرجات!{tc}{start + 1}:{tc}{last_row}),0)"),
        ("الأدنى", f"=IFERROR(MIN(الدرجات!{tc}{start + 1}:{tc}{last_row}),0)"),
        ("الانحراف المعياري",
         f"=IFERROR(STDEV(الدرجات!{tc}{start + 1}:{tc}{last_row}),0)"),
        ("تحت 50%",
         f'=COUNTIF(الدرجات!{tc}{start + 1}:{tc}{last_row},"<{max_points / 2}")'),
    ]
    stats["A1"], stats["B1"] = "المؤشر", "القيمة"
    for cell in ("A1", "B1"):
        stats[cell].font = header_font
        stats[cell].fill = header_fill
    for i, (label, formula) in enumerate(rows, start=2):
        stats.cell(row=i, column=1, value=label).font = Font(name="Arial", size=10)
        stats.cell(row=i, column=2, value=formula).number_format = "0.00"
    stats.column_dimensions["A"].width = 20
    stats.column_dimensions["B"].width = 14

    out = work_dir / "grades_draft.xlsx"
    wb.save(out)
    return out


def main() -> None:
    # مثل cli.py: الإخراج عربي دائماً، حتى على كونسول ويندوز القديم (cp1252).
    sys.stdout.reconfigure(encoding="utf-8")

    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)

    work_dir = Path(sys.argv[1])
    data = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))

    criteria = data["criteria"]
    total_possible = sum(c["points"] for c in criteria)
    max_points = data.get("max_points", total_possible)

    if abs(total_possible - max_points) > 0.01:
        print(f"⚠️  مجموع المعايير {total_possible} ≠ العلامة الكاملة {max_points}")

    out = write_grades(work_dir, data)
    print(f"✓ {out}")
    print(f"  {len(data['grades'])} طالب  |  {len(criteria)} معيار  |  "
          f"العلامة الكاملة {max_points}")
    print("  راجع الملف قبل الرفع.")


def _load_names(roster_path: Path) -> dict:
    if not roster_path.exists():
        return {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(roster_path, data_only=True)
        ws = wb["التسليمات"]
        names = {}
        for row in ws.iter_rows(min_row=6, max_col=2, values_only=True):
            if row[0]:
                names[str(row[0])] = row[1] or ""
        return names
    except Exception:
        return {}


if __name__ == "__main__":
    main()
