"""الأمر status: تقرير متابعة لكل طالب عبر كل واجبات المساق.

`compute_status()` يبني البيانات (matrix / rows / summary) بدون أي I/O للملفات
أو طباعة — شاشة تقرير المتابعة بتستهلكها مباشرة. `write_status_xlsx()` يكتب نفس
ملف الإكسل القديم من تلك البيانات. `status()` يربط الاثنين ويطبع كما قبل.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import api
from .pull import SUBMITTED_STATES, build_student_index

_GOOD = "جيد"
_AT_RISK = "⚠️ متابعة"


def compute_status(classroom, cfg: dict, course_id: str,
                   risk_threshold: float = 0.6) -> dict:
    """{works, matrix[uid][workId], rows[...], summary} — pure computation."""
    works = [w for w in api.list_coursework(classroom, course_id)
             if w.get("workType") == "ASSIGNMENT"]
    if not works:
        raise SystemExit("✗ ما في واجبات في هذا المساق.")

    students = build_student_index(classroom, course_id, cfg["student_id_pattern"])

    matrix: dict[str, dict[str, dict]] = {uid: {} for uid in students}
    for work in works:
        for sub in api.list_submissions(classroom, course_id, work["id"]):
            uid = sub["userId"]
            if uid not in matrix:
                continue
            matrix[uid][work["id"]] = {
                "submitted": sub.get("state") in SUBMITTED_STATES,
                "late": bool(sub.get("late")),
                "grade": sub.get("assignedGrade"),
            }

    ordered = sorted(
        students.items(),
        key=lambda kv: (kv[1]["student_id"] is None, kv[1]["student_id"] or ""),
    )

    rows: list[dict] = []
    for uid, info in ordered:
        done = late_count = 0
        grades: list[float] = []
        cells: list[str] = []
        for work in works:
            rec = matrix[uid].get(work["id"])
            if rec and rec["submitted"]:
                done += 1
                if rec["late"]:
                    late_count += 1
                    cells.append("late")
                else:
                    cells.append("ok")
                if rec["grade"] is not None:
                    grades.append(rec["grade"])
            else:
                cells.append("missing")

        ratio = done / len(works) if works else 0
        avg = sum(grades) / len(grades) if grades else None
        rows.append({
            "user_id": uid,
            "student_id": info["student_id"] or "",
            "name": info["full_name"],
            "cells": cells,
            "done": done,
            "ratio": ratio,
            "late": late_count,
            "avg": round(avg, 2) if avg is not None else None,
            "status": _AT_RISK if ratio < risk_threshold else _GOOD,
        })

    at_risk = sorted((r for r in rows if r["ratio"] < risk_threshold),
                     key=lambda r: r["ratio"])
    summary = {
        "course_id": course_id,
        "total_students": len(students),
        "total_works": len(works),
        "threshold": risk_threshold,
        "at_risk": at_risk,
    }
    return {"works": works, "matrix": matrix, "rows": rows, "summary": summary}


def write_status_xlsx(computed: dict, course_key: str, out_dir: Path) -> Path:
    """يكتب `_status_YYYYMMDD.xlsx` من مخرجات `compute_status` (نفس التنسيق القديم)."""
    works = computed["works"]
    rows = computed["rows"]
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"_status_{datetime.now():%Y%m%d}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "المتابعة"
    ws.sheet_view.rightToLeft = True

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
    header_fill = PatternFill("solid", fgColor="2F5597")
    ok_fill = PatternFill("solid", fgColor="C6EFCE")
    late_fill = PatternFill("solid", fgColor="FFEB9C")
    miss_fill = PatternFill("solid", fgColor="FFC7CE")
    risk_fill = PatternFill("solid", fgColor="FF9999")

    ws["A1"] = f"تقرير متابعة — {course_key}"
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws["A2"] = ("✓ سلّم   |   ⏰ متأخر   |   ✗ لم يسلّم        "
                f"أُنشئ: {datetime.now():%Y-%m-%d %H:%M}")
    ws["A2"].font = Font(name="Arial", italic=True, size=9)

    start = 4
    headers = ["الرقم الجامعي", "الاسم"] + [w.get("title", "")[:18] for w in works] \
        + ["نسبة التسليم", "متأخر", "متوسط الدرجة", "حالة"]
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=start, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True,
                                   textRotation=45 if 2 < col <= 2 + len(works) else 0)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    for col in range(3, 3 + len(works)):
        ws.column_dimensions[get_column_letter(col)].width = 6
    for col in range(3 + len(works), 3 + len(works) + 4):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.row_dimensions[start].height = 70

    _mark = {"ok": ("✓", ok_fill), "late": ("⏰", late_fill), "missing": ("✗", miss_fill)}
    for offset, row in enumerate(rows, start=1):
        r = start + offset
        ws.cell(row=r, column=1, value=row["student_id"])
        ws.cell(row=r, column=2, value=row["name"])

        for i, kind in enumerate(row["cells"]):
            cell = ws.cell(row=r, column=3 + i)
            cell.alignment = Alignment(horizontal="center")
            cell.value, cell.fill = _mark[kind]

        base = 3 + len(works)
        ws.cell(row=r, column=base, value=row["ratio"]).number_format = "0%"
        ws.cell(row=r, column=base + 1, value=row["late"])
        ws.cell(row=r, column=base + 2,
                value=row["avg"] if row["avg"] is not None else "—")

        flag = ws.cell(row=r, column=base + 3, value=row["status"])
        if row["status"] == _AT_RISK:
            flag.fill = risk_fill

        for col in range(1, base + 4):
            cell = ws.cell(row=r, column=col)
            if not cell.font.bold:
                cell.font = Font(name="Arial", size=9)

    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.auto_filter.ref = (
        f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"
    )
    wb.save(path)
    return path


def status(classroom, cfg: dict, course_key: str, course_id: str,
           risk_threshold: float = 0.6) -> Path:
    computed = compute_status(classroom, cfg, course_id, risk_threshold)
    s = computed["summary"]
    print(f"📊 {s['total_students']} طالب  ×  {s['total_works']} واجب")
    for work in computed["works"]:
        print(f"   ← {work.get('title')}")

    path = write_status_xlsx(computed, course_key, Path(cfg["output_dir"]) / course_key)
    print(f"\n✅ {path}")

    at_risk = s["at_risk"]
    if at_risk:
        print(f"\n⚠️  {len(at_risk)} طالب تحت عتبة {risk_threshold:.0%}:")
        for row in at_risk[:10]:
            print(f"     {row['student_id'] or '?':<12} {row['name']:<28} {row['ratio']:.0%}")
    return path
