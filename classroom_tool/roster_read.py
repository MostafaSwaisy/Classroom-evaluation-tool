"""Read the per-assignment `_roster.xlsx` and `_missing.txt` for the UI.

Read-only: these files are produced by `pull` and are never written back from
the GUI (CLAUDE.md guardrail). `read_roster` returns the sheet as a list of
plain dicts keyed by stable English names.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

_SHEET = "التسليمات"
_ID_HEADER = "الرقم الجامعي"

# Arabic header label -> dict key. Order here is the natural column order.
_FIELD_BY_HEADER: dict[str, str] = {
    "الرقم الجامعي": "student_id",
    "الاسم": "name",
    "الإيميل": "email",
    "الحالة": "state",
    "متأخر": "late",
    "وقت التسليم": "turned_in",
    "عدد الملفات": "n_files",
    "الملفات": "files",
    "روابط": "links",
    "الدرجة الحالية": "current_grade",
    "الدرجة": "grade",
    "ملاحظات": "notes",
}

#: the ten fields the roster screen requires (spec §5.9).
SPEC_FIELDS: tuple[str, ...] = (
    "student_id", "name", "email", "state", "late",
    "turned_in", "n_files", "files", "links", "current_grade",
)

_TEXT_FIELDS = {"student_id", "name", "email", "state", "turned_in", "files", "links", "notes"}


def read_roster(path: str | Path) -> list[dict]:
    """Rows of `_roster.xlsx` as dicts. Missing file -> `[]`."""
    path = Path(path)
    if not path.exists():
        return []

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[_SHEET] if _SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        col_key = _header_map(rows)

        out: list[dict] = []
        for raw in rows:
            if _blank(raw):
                continue
            rec = dict.fromkeys(_FIELD_BY_HEADER.values())
            for idx, key in col_key.items():
                rec[key] = raw[idx] if idx < len(raw) else None
            out.append(_normalise(rec))
        return out
    finally:
        wb.close()


def read_missing(path: str | Path) -> list[str]:
    """Student names from `_missing.txt` (header lines dropped). Missing file -> `[]`."""
    path = Path(path)
    if not path.exists():
        return []
    names: list[str] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        s = line.strip()
        if not s or s.startswith(("لم يسلّموا", "العدد:")):
            continue
        names.append(s)
    return names


# --- internals --------------------------------------------------------------

def _header_map(rows) -> dict[int, str]:
    for raw in rows:
        if raw and raw[0] and str(raw[0]).strip() == _ID_HEADER:
            return {
                i: _FIELD_BY_HEADER[str(h).strip()]
                for i, h in enumerate(raw)
                if h is not None and str(h).strip() in _FIELD_BY_HEADER
            }
    raise ValueError(f"header row (starting with {_ID_HEADER!r}) not found in roster")


def _blank(raw) -> bool:
    return raw is None or all(c is None or str(c).strip() == "" for c in raw)


def _normalise(rec: dict) -> dict:
    rec["late"] = str(rec.get("late") or "").strip() == "نعم"
    n = rec.get("n_files")
    rec["n_files"] = int(n) if isinstance(n, (int, float)) else 0
    for key in _TEXT_FIELDS:
        v = rec.get(key)
        rec[key] = "" if v is None else str(v).strip()
    # grade / current_grade: keep the raw value (None, number, or text)
    return rec
